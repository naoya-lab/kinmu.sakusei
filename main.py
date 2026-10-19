"""Monthly radiation allowance from one staff member's PDF schedule."""
import argparse, calendar, csv, json, re
from pathlib import Path
import fitz
from openpyxl import load_workbook

CONFIG = Path(__file__).with_name("local_config.json")
EXCLUDE = ("週", "休", "年休", "夏季", "子看", "服忌", "代", "欠", "出張")
TITLE = re.compile(r"令和\s*(\d+)\s*年\s*(\d+)\s*月\s*放射線部勤務表")

def extract(pdf, name):
    with fitz.open(pdf) as doc:
        if len(doc) != 1: raise ValueError("1-page PDF is required")
        spans = [s for b in doc[0].get_text("dict")["blocks"] for line in b.get("lines", []) for s in line["spans"]]
    titles = [m for s in spans if (m := TITLE.search(s["text"]))]
    if len(titles) != 1: raise ValueError("年月を判定できません")
    year, month = int(titles[0][1]) + 2018, int(titles[0][2])
    names = [s for s in spans if re.sub(r"\s", "", s["text"]) == re.sub(r"\s", "", name)]
    if len(names) != 1: raise ValueError("氏名が一意に見つかりません")
    cy = sum((names[0]["bbox"][1], names[0]["bbox"][3])) / 2
    hdr = [s for s in spans if s["text"].strip().isascii() and s["text"].strip().isdigit()
           and 1 <= int(s["text"].strip()) <= 31 and 70 < s["bbox"][1] < cy]
    if not hdr: raise ValueError("日付が見つかりません")
    y0 = min(s["bbox"][1] for s in hdr)
    hdr = sorted([s for s in hdr if abs(s["bbox"][1] - y0) < 4], key=lambda s: s["bbox"][0])
    days = calendar.monthrange(year, month)[1]
    candidates = [i for i in range(len(hdr) - days + 1)
                  if all(hdr[i + k - 1]["text"].strip() == str(k) for k in range(1, days + 1))]
    if len(candidates) != 1: raise ValueError("対象月の日付配置が不明です")
    head = hdr[candidates[0]:candidates[0] + days]
    centers = [(s["bbox"][0] + s["bbox"][2]) / 2 for s in head]
    pitch = sorted([centers[i + 1] - centers[i] for i in range(days - 1)])[days // 2 - 1]
    if not 15 < pitch < 60: raise ValueError("列間隔が異常です")
    rows = []
    for d in range(1, days + 1):
        hits = []
        for s in spans:
            x = sum((s["bbox"][0], s["bbox"][2])) / 2
            y = sum((s["bbox"][1], s["bbox"][3])) / 2
            if abs(y - cy) < 8 and abs(x - centers[d - 1]) <= pitch * .48:
                hits.append(s)
        hits.sort(key=lambda x: x["bbox"][0])
        value = "".join(s["text"].strip() for s in hits)
        red = any((s["color"] >> 16) & 255 >= 170 and (s["color"] >> 8) & 255 <= 110
                  and s["color"] & 255 <= 110 for s in hits)
        off = any(v in value for v in EXCLUDE)
        rows.append((d, value, red, off, bool(value) and not red and not off))
    return year, month, rows

def main():
    p = argparse.ArgumentParser()
    p.add_argument("pdf", type=Path)
    p.add_argument("template", type=Path)
    p.add_argument("--name")
    p.add_argument("--output", type=Path, default=Path("放射線取扱手当_記入済.xlsx"))
    p.add_argument("--review", type=Path, default=Path("判定確認.csv"))
    a = p.parse_args()
    if a.name:
        name = a.name
        CONFIG.write_text(json.dumps({"name": name}, ensure_ascii=False), encoding="utf-8")
    elif CONFIG.exists():
        name = json.loads(CONFIG.read_text(encoding="utf-8"))["name"]
    else:
        name = input("対象職員の氏名（初回のみ）: ").strip()
        if not name: raise ValueError("氏名が空です")
        CONFIG.write_text(json.dumps({"name": name}, ensure_ascii=False), encoding="utf-8")
    year, month, rows = extract(a.pdf, name)
    wb = load_workbook(a.template)
    if "様式" not in wb.sheetnames: raise ValueError("様式シートがありません")
    ws = wb["様式"]
    for day in range(1, 32): ws.cell(day + 4, 8).value = None
    ws["A1"] = f"　　　　　　　　　　　　　　{month} 月　放　射　線　取　扱　作　業　等　手　当"
    for day, value, red, off, mark in rows:
        if mark: ws.cell(day + 4, 8).value = "〇"
    ws["H36"] = '=COUNTIF(H5:I35,"〇")'
    wb.save(a.output)
    with a.review.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["年月", "日", "勤務区分", "赤字あり", "休日等", "丸印"])
        for day, value, red, off, mark in rows:
            w.writerow([f"{year}-{month:02}", day, value, "はい" if red else "", "はい" if off else "", "〇" if mark else ""])
    print(f"{year}/{month}: 仮判定 {sum(r[4] for r in rows)} 日。提出前に判定確認.csvを確認してください。")
if __name__ == "__main__": main()
